"""Reading and checking FDR rows from a CSV or Excel file. Nothing here touches the database."""

import csv
import io
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from typing import Any

from .schemas import clean_decimal

MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 500

COLUMNS = ["reference_no", "bank_name", "amount", "issue_date", "maturity_date"]

# Header names people commonly use in their own sheets, mapped to our column names.
ALIASES = {
    "fdr_no": "reference_no",
    "fdr_number": "reference_no",
    "fdr": "reference_no",
    "reference": "reference_no",
    "bank": "bank_name",
    "amount_inr": "amount",
    "amount_rs": "amount",
    "issue": "issue_date",
    "maturity": "maturity_date",
}

DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%d-%m-%Y", "%d.%m.%Y")


class ImportFileError(ValueError):
    """The file itself cannot be used (wrong type, too big, too many rows, unreadable)."""


def _normal_header(name: Any) -> str:
    text = str(name or "").strip().lower()
    return "_".join(text.replace("-", " ").split())


def _cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return str(value).strip()


def read_table(data: bytes, filename: str) -> list[dict[str, str]]:
    """Return one dict per data row, keyed by our column names. Blank rows are skipped."""
    if len(data) > MAX_IMPORT_BYTES:
        raise ImportFileError("File is larger than 5 MB")

    name = (filename or "").lower()
    if name.endswith(".xlsx"):
        raw_rows = _xlsx_rows(data)
    elif name.endswith(".csv"):
        raw_rows = _csv_rows(data)
    else:
        raise ImportFileError("Upload a .csv or .xlsx file")

    if not raw_rows:
        raise ImportFileError("The file is empty")

    headers = [_normal_header(h) for h in raw_rows[0]]
    headers = [ALIASES.get(h, h) for h in headers]
    missing = [c for c in COLUMNS if c not in headers]
    if missing:
        raise ImportFileError("Missing column(s): " + ", ".join(missing))

    rows: list[dict[str, str]] = []
    for raw in raw_rows[1:]:
        cells = [_cell_text(v) for v in raw]
        if not any(cells):
            continue
        row = {}
        for index, header in enumerate(headers):
            if header in COLUMNS:
                row[header] = cells[index] if index < len(cells) else ""
        rows.append(row)

    if len(rows) > MAX_IMPORT_ROWS:
        raise ImportFileError(f"A file can hold at most {MAX_IMPORT_ROWS} FDRs")
    if not rows:
        raise ImportFileError("The file has a header but no FDR rows")
    return rows


def _csv_rows(data: bytes) -> list[list[str]]:
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    return list(csv.reader(io.StringIO(text)))


def _xlsx_rows(data: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl raises several unrelated error types
        raise ImportFileError("This Excel file could not be read") from exc
    sheet = workbook.worksheets[0]
    return [list(r) for r in sheet.iter_rows(values_only=True)]


def parse_date(value: str) -> date | None:
    for fmt in DATE_FORMATS:
        try:
            return datetime.strptime(value, fmt).date()
        except ValueError:
            continue
    return None


def check_rows(rows: list[dict[str, str]], existing_refs: set[str]) -> list[dict[str, Any]]:
    """Check each row. Returns one result per row with the cleaned values and any errors."""
    results: list[dict[str, Any]] = []
    first_line_for_ref: dict[str, int] = {}

    for index, row in enumerate(rows):
        line = index + 2  # line 1 is the header
        errors: list[str] = []
        ref = row["reference_no"].strip()
        bank = row["bank_name"].strip()

        if not ref:
            errors.append("FDR number is missing")
        elif len(ref) > 100:
            errors.append("FDR number is longer than 100 characters")
        if not bank:
            errors.append("Bank name is missing")
        elif len(bank) > 255:
            errors.append("Bank name is longer than 255 characters")

        amount = None
        try:
            amount = Decimal(str(clean_decimal(row["amount"]) or "").strip())
            if amount <= 0:
                errors.append("Amount must be greater than zero")
        except (InvalidOperation, ValueError, TypeError):
            errors.append("Amount is not a valid number")

        issue = parse_date(row["issue_date"])
        maturity = parse_date(row["maturity_date"])
        if issue is None:
            errors.append("Issue date is not recognised (use YYYY-MM-DD or DD/MM/YYYY)")
        if maturity is None:
            errors.append("Maturity date is not recognised (use YYYY-MM-DD or DD/MM/YYYY)")
        if issue and maturity and maturity <= issue:
            errors.append("Maturity date must be after the issue date")

        if ref:
            if ref.lower() in existing_refs:
                errors.append("This FDR number is already in the portal")
            if ref.lower() in first_line_for_ref:
                errors.append(f"This FDR number also appears on line {first_line_for_ref[ref.lower()]} of the file")
            else:
                first_line_for_ref[ref.lower()] = line

        results.append(
            {
                "line": line,
                "data": {
                    "reference_no": ref,
                    "bank_name": bank,
                    "amount": str(amount) if amount is not None else row["amount"],
                    "issue_date": issue.isoformat() if issue else row["issue_date"],
                    "maturity_date": maturity.isoformat() if maturity else row["maturity_date"],
                },
                "errors": errors,
            }
        )
    return results


def template_csv() -> str:
    return (
        ",".join(COLUMNS)
        + "\n"
        + "FDR-2026-0001,State Bank of India,500000.00,2026-04-01,2027-04-01\n"
        + "FDR-2026-0002,Canara Bank,250000.50,15/05/2026,15/05/2027\n"
    )
